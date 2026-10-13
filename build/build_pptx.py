"""Dựng bộ slide 12 trang từ ảnh và chú giải do notebook sinh ra."""
import sys
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from data import ROOT, FigNote, SummaryData, aspect, fmt_int, fmt_num, load_note, load_summary, shorten

OUT_PPTX = ROOT / "deliverables" / "slides.pptx"
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.5
TITLE_TOP, TITLE_H = 0.3, 0.95
BODY_TOP = 1.55
BULLET_LIMIT = 230
TITLE_PT, BODY_PT, SMALL_PT = 26, 18, 11
SINGLE_IMG_W = 8.6
PAIR_GAP = 0.3
PAIR_BODY_PT = 17
JOIN_SKIP = ("f6_3", "f5_2")
INK = RGBColor(0x1F, 0x29, 0x37)
ACCENT = RGBColor(0x00, 0x72, 0xB2)
ACCENT2 = RGBColor(0xD5, 0x5E, 0x00)
GREY = RGBColor(0x66, 0x66, 0x66)
BLANK_LAYOUT = 6
TEAM = "Nhóm 8"
COURSE = "CSE702060 – Trực quan hóa dữ liệu"
MEMBERS = "Lê Đức Duy · Vũ Viết Huy · Lê Mạnh Hùng · Nguyễn Bá Huy"
QUESTIONS = (
    "Q1 Ai là người trả lời khảo sát, mẫu có lệch không?", "Q2 Công nghệ nào đang được dùng nhiều nhất?",
    "Q3 Dùng nhiều có được ngưỡng mộ và muốn dùng nhiều không?", "Q4 Các ngôn ngữ đi cùng nhau thế nào, ai muốn dùng gì?",
    "Q5 Phân phối lương theo vai trò và kinh nghiệm?", "Q6 Công nghệ nào đi cùng mức lương cao, sau khi tách nước và kinh nghiệm?",
    "Q7 Lương khác nhau thế nào giữa các nước, Việt Nam đứng đâu?", "Q8 So sánh các vai trò.",
)


@dataclass
class SlideSpec:
    """Nội dung một slide biểu đồ: tiêu đề, danh sách mã biểu đồ, ý chính và ghi chú trình bày."""
    title: str
    codes: list
    bullets: list = field(default_factory=list)
    speaker: str = ""


def add_text(slide, text, left, top, width, height, size, bold=False, color=INK, align=PP_ALIGN.LEFT):
    """Thêm một hộp văn bản một đoạn."""
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    para = frame.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.size, run.font.bold, run.font.color.rgb = Pt(size), bold, color
    return box


def add_bullets(slide, items, left, top, width, height, size=BODY_PT):
    """Thêm danh sách ý chính, mỗi ý một đoạn."""
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    for idx, text in enumerate(items):
        para = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        para.space_after = Pt(8)
        run = para.add_run()
        run.text = f"• {text}"
        run.font.size, run.font.color.rgb = Pt(size), INK


def add_frame(slide, title, number):
    """Thêm dải màu nhấn, tiêu đề nêu phát hiện và số trang."""
    bar = slide.shapes.add_shape(1, Inches(0), Inches(TITLE_TOP), Inches(0.15), Inches(TITLE_H - 0.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    add_text(slide, title, MARGIN, TITLE_TOP - 0.05, SLIDE_W - 2 * MARGIN, TITLE_H, TITLE_PT, bold=True)
    add_text(slide, f"{TEAM} · {COURSE} · {number}", MARGIN, SLIDE_H - 0.4, SLIDE_W - 2 * MARGIN, 0.3, SMALL_PT, color=GREY, align=PP_ALIGN.RIGHT)


def place_images(slide, notes):
    """Đặt 1 ảnh (trái) hoặc 2 ảnh cạnh nhau; trả về vùng còn lại cho chữ (left, top, width, height)."""
    usable_w = SLIDE_W - 2 * MARGIN
    if len(notes) == 1:
        slide.shapes.add_picture(str(notes[0].cropped), Inches(MARGIN), Inches(BODY_TOP), width=Inches(SINGLE_IMG_W))
        return MARGIN + SINGLE_IMG_W + PAIR_GAP, BODY_TOP, usable_w - SINGLE_IMG_W - PAIR_GAP, SLIDE_H - BODY_TOP - 0.6
    width = (usable_w - PAIR_GAP) / 2
    heights = []
    for idx, note in enumerate(notes):
        slide.shapes.add_picture(str(note.cropped), Inches(MARGIN + idx * (width + PAIR_GAP)), Inches(BODY_TOP), width=Inches(width))
        heights.append(width * aspect(note.cropped))
    top = BODY_TOP + max(heights) + 0.2
    return MARGIN, top, usable_w, SLIDE_H - top - 0.5


def figure_slide(prs, spec: SlideSpec, number: int):
    """Dựng một slide biểu đồ."""
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT])
    add_frame(slide, spec.title, number)
    notes = [load_note(code) for code in spec.codes]
    left, top, width, height = place_images(slide, notes)
    add_bullets(slide, spec.bullets, left, top, width, height, size=BODY_PT if len(notes) == 1 else PAIR_BODY_PT)
    slide.notes_slide.notes_text_frame.text = spec.speaker


def title_slide(prs, s: SummaryData):
    """Slide 1: tiêu đề, nhóm và thành viên."""
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT])
    add_text(slide, "Phân tích xu hướng công nghệ và mức lương lập trình viên", 0.8, 2.0, 11.7, 1.6, 40, bold=True, color=ACCENT)
    add_text(slide, "Qua khảo sát Stack Overflow Developer Survey 2025", 0.8, 3.6, 11.7, 0.6, 24)
    add_text(slide, f"{TEAM} · {COURSE}", 0.8, 4.6, 11.7, 0.5, 18, color=GREY)
    add_text(slide, MEMBERS, 0.8, 5.1, 11.7, 0.5, 16, color=GREY)
    add_text(slide, f"Dữ liệu: {fmt_int(s.n_all)} phản hồi · Phenikaa School of Computing · 10/2026", 0.8, 6.3, 11.7, 0.4, SMALL_PT + 1, color=GREY)


def context_slide(prs, s: SummaryData, number: int):
    """Slide 2: bối cảnh, câu hỏi và đối tượng."""
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT])
    add_frame(slide, "Công nghệ nào nên học và kỳ vọng mức lương ra sao? Dữ liệu của 49 nghìn người trả lời", number)
    add_bullets(slide, [
        f"Dữ liệu: Stack Overflow Developer Survey 2025, {fmt_int(s.n_all)} người trả lời; {s.pct_pro:.0f}% là lập trình viên chuyên nghiệp.".replace(".0", ""),
        f"Phân tích lương: {fmt_int(s.n_sal)} lập trình viên chuyên nghiệp có lương hợp lệ (1.000–500.000 USD); dùng median vì lương lệch phải.",
        "Đối tượng: sinh viên và lập trình viên mới ở Việt Nam, cần chọn công nghệ học và đặt kỳ vọng lương.",
        "Xu hướng: dữ liệu chỉ có năm 2025, nên dùng ba chỉ số thay thế: đang dùng, muốn dùng, được ngưỡng mộ.",
    ], MARGIN, BODY_TOP, 6.2, 5.2, size=18)
    add_text(slide, "Tám câu hỏi chính (câu thứ chín về kỹ năng và AI nằm trong báo cáo)", 7.2, BODY_TOP, 5.6, 0.8, 16, bold=True, color=ACCENT)
    add_bullets(slide, list(QUESTIONS), 7.2, BODY_TOP + 1.0, 5.6, 4.5, size=15)
    slide.notes_slide.notes_text_frame.text = "Nêu bối cảnh, đối tượng, và cách đọc: median, n, USD danh nghĩa."


def kpi_card(slide, left, value, label, color):
    """Một thẻ số liệu lớn kèm nhãn."""
    add_text(slide, value, left, 1.5, 3.9, 1.0, 44, bold=True, color=color, align=PP_ALIGN.CENTER)
    add_text(slide, label, left, 2.5, 3.9, 0.8, 14, color=GREY, align=PP_ALIGN.CENTER)


def message_slide(prs, s: SummaryData, number: int):
    """Slide 11: thông điệp chính và khuyến nghị."""
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT])
    add_frame(slide, "Quốc gia và kinh nghiệm đi cùng chênh lệch lương nhiều hơn ngôn ngữ lập trình", number)
    pct = lambda v: f"{fmt_num(v * 100, 0)}%"
    kpi_card(slide, MARGIN, pct(s.r2_country), "biến thiên log(lương) do quốc gia giải thích", ACCENT)
    kpi_card(slide, MARGIN + 4.2, pct(s.r2_country_exp), "khi thêm kinh nghiệm làm việc", ACCENT)
    kpi_card(slide, MARGIN + 8.4, f"+{fmt_num((s.r2_full - s.r2_country_exp) * 100, 1)} điểm %", "chỉ thêm khi cộng 15 ngôn ngữ phổ biến", ACCENT2)
    add_bullets(slide, s.recommendations, MARGIN, 3.5, SLIDE_W - 2 * MARGIN, 3.6, size=16)
    slide.notes_slide.notes_text_frame.text = s.message


def limits_slide(prs, s: SummaryData, number: int):
    """Slide 12: hạn chế và Q&A."""
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT])
    add_frame(slide, "Kết quả mang tính mô tả của mẫu khảo sát, cần đọc cùng các hạn chế", number)
    add_bullets(slide, s.limitations, MARGIN, BODY_TOP, SLIDE_W - 2 * MARGIN, 4.2, size=18)
    add_text(slide, "Cảm ơn thầy cô và các bạn đã lắng nghe. Q&A", MARGIN, 6.2, SLIDE_W - 2 * MARGIN, 0.6, 24, bold=True, color=ACCENT)


def spec_from(codes: list, title_code: str = None, extra: str = "", join_titles: bool = False) -> SlideSpec:
    """Dựng SlideSpec từ chú giải: tiêu đề là phát hiện, ý chính là câu đầu của nhận xét và insight."""
    notes = [load_note(code) for code in codes]
    head = load_note(title_code or codes[0])
    title = "; ".join([load_note(codes[0]).title] + [load_note(c).title for c in codes[1:]]) if join_titles else None
    bullets = [shorten(n.observation, BULLET_LIMIT) for n in notes] + [shorten(notes[-1].insight, BULLET_LIMIT)]
    if extra:
        bullets.insert(0, extra)
    speaker = "\n\n".join(f"{n.title}. {n.observation} {n.insight}" for n in notes)
    return SlideSpec(title or head.title, codes, bullets, speaker)


def build(prs: Presentation, s: SummaryData) -> None:
    """Dựng đủ 12 slide theo thứ tự câu chuyện: rộng đến hẹp."""
    title_slide(prs, s)
    context_slide(prs, s, 2)
    plan = (["f1_3"], ["f2_1", "f2_2"], ["f3_1"], ["f4_1", "f4_2"], ["f5_2", "f5_3"], ["f6_3", "f6_4"], ["f7_1"], ["f8_1"])
    for number, codes in enumerate(plan, start=3):
        figure_slide(prs, spec_from(codes, join_titles=len(codes) == 2 and codes[0] not in JOIN_SKIP), number)
    message_slide(prs, s, 11)
    limits_slide(prs, s, 12)


def main() -> None:
    """Điểm vào: tạo và lưu bộ slide."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    build(prs, load_summary())
    OUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT_PPTX)
    print(f"Đã lưu {OUT_PPTX} ({len(prs.slides)} slide)")


if __name__ == "__main__":
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    main()
